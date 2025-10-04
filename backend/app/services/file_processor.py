import os
import tempfile
import asyncio
from typing import List, Tuple, Optional
from PIL import Image
import fitz  # PyMuPDF
from pptx import Presentation
import docx
import pdfkit
import logging
from concurrent.futures import ThreadPoolExecutor

from app.core.config import settings
from app.services.storage_service import StorageService

logger = logging.getLogger(__name__)

class FileProcessor:
    def __init__(self):
        self.storage_service = StorageService()
        self.executor = ThreadPoolExecutor(max_workers=4)

    async def process_file(self, task_id: str, file_path: str, file_type: str) -> List[Tuple[int, str]]:
        """处理文件并返回页面URL列表"""
        try:
            with tempfile.TemporaryDirectory() as temp_dir:
                logger.info(f"Processing file for task {task_id}, type: {file_type}")
                
                # 根据文件类型处理
                if file_type == "application/pdf":
                    pages = await self._process_pdf(file_path, temp_dir)
                elif file_type in ["application/vnd.ms-powerpoint", 
                                   "application/vnd.openxmlformats-officedocument.presentationml.presentation"]:
                    pages = await self._process_ppt(file_path, temp_dir)
                elif file_type in ["application/msword",
                                   "application/vnd.openxmlformats-officedocument.wordprocessingml.document"]:
                    pages = await self._process_word(file_path, temp_dir)
                elif file_type in ["image/jpeg", "image/png"]:
                    pages = await self._process_image(file_path, temp_dir)
                else:
                    raise ValueError(f"Unsupported file type: {file_type}")
                
                # 上传页面图片
                page_urls = []
                for page_num, page_path in pages:
                    object_key = f"flipbooks/{task_id}/page_{page_num}.png"
                    url = await self.storage_service.upload_file(
                        page_path, 
                        object_key, 
                        "image/png"
                    )
                    page_urls.append((page_num, url))
                
                logger.info(f"Processed {len(page_urls)} pages for task {task_id}")
                return page_urls
                
        except Exception as e:
            logger.error(f"Error processing file for task {task_id}: {e}")
            raise

    async def _process_pdf(self, pdf_path: str, temp_dir: str) -> List[Tuple[int, str]]:
        """处理PDF文件"""
        def process_pdf_sync():
            doc = fitz.open(pdf_path)
            pages = []
            
            for page_num in range(doc.page_count):
                page = doc[page_num]
                # 设置高分辨率
                mat = fitz.Matrix(settings.PDF_DPI / 72, settings.PDF_DPI / 72)
                pix = page.get_pixmap(matrix=mat)
                
                # 保存为PNG
                img_path = os.path.join(temp_dir, f"page_{page_num + 1}.png")
                pix.save(img_path)
                
                # 优化图片尺寸
                optimized_path =  self._optimize_image_sync(img_path, temp_dir, page_num + 1)
                pages.append((page_num + 1, optimized_path))
            
            doc.close()
            return pages
        
        loop = asyncio.get_event_loop()
        return await loop.run_in_executor(self.executor, process_pdf_sync)

    async def _process_ppt(self, ppt_path: str, temp_dir: str) -> List[Tuple[int, str]]:
        """处理PPT文件"""
        def process_ppt_sync():
            try:
                # 使用LibreOffice将PPT转为PDF，然后转图片
                pdf_path = os.path.join(temp_dir, "temp.pdf")
                
                # 这里需要LibreOffice headless模式
                # 简化实现：创建占位图片
                prs = Presentation(ppt_path)
                pages = []
                
                for i, slide in enumerate(prs.slides):
                    # 创建占位图片（实际项目中需要真正渲染幻灯片）
                    img = Image.new('RGB', (1920, 1080), 'white')
                    img_path = os.path.join(temp_dir, f"slide_{i + 1}.png")
                    img.save(img_path)
                    pages.append((i + 1, img_path))
                
                return pages
            except Exception as e:
                logger.error(f"PPT processing error: {e}")
                return []
        
        loop = asyncio.get_event_loop()
        return await loop.run_in_executor(self.executor, process_ppt_sync)

    async def _process_word(self, word_path: str, temp_dir: str) -> List[Tuple[int, str]]:
        """处理Word文件"""
        def process_word_sync():
            try:
                # 读取Word文档内容
                doc = docx.Document(word_path)
                text_content = "\n".join([paragraph.text for paragraph in doc.paragraphs])
                safe_text = text_content.replace('\n', '<br>')
                # 转为HTML
                html_content = f"""
                <html>
                <head>
                    <style>
                        body {{ font-family: Arial, sans-serif; margin: 40px; line-height: 1.6; }}
                        p {{ margin-bottom: 10px; }}
                    </style>
                </head>
                <body>
                    {safe_text}
                </body>
                </html>
                """
                
                # HTML转PDF
                pdf_path = os.path.join(temp_dir, "temp.pdf")
                pdfkit.from_string(html_content, pdf_path)
                
                # PDF转图片
                return self._pdf_to_images_sync(pdf_path, temp_dir)
                
            except Exception as e:
                logger.error(f"Word processing error: {e}")
                return []
        
        loop = asyncio.get_event_loop()
        return await loop.run_in_executor(self.executor, process_word_sync)

    async def _process_image(self, image_path: str, temp_dir: str) -> List[Tuple[int, str]]:
        """处理单张图片"""
        def process_image_sync():
            try:
                img = Image.open(image_path)
                
                # 调整尺寸
                img.thumbnail((settings.IMAGE_MAX_WIDTH, settings.IMAGE_MAX_HEIGHT), Image.Resampling.LANCZOS)
                
                # 保存优化后的图片
                output_path = os.path.join(temp_dir, "page_1.png")
                img.save(output_path, "PNG", quality=settings.IMAGE_QUALITY, optimize=True)
                
                return [(1, output_path)]
            except Exception as e:
                logger.error(f"Image processing error: {e}")
                return []
        
        loop = asyncio.get_event_loop()
        return await loop.run_in_executor(self.executor, process_image_sync)

    def _pdf_to_images_sync(self, pdf_path: str, temp_dir: str) -> List[Tuple[int, str]]:
        """同步版本PDF转图片"""
        doc = fitz.open(pdf_path)
        pages = []
        
        for page_num in range(doc.page_count):
            page = doc[page_num]
            mat = fitz.Matrix(2.0, 2.0)
            pix = page.get_pixmap(matrix=mat)
            
            img_path = os.path.join(temp_dir, f"page_{page_num + 1}.png")
            pix.save(img_path)
            pages.append((page_num + 1, img_path))
        
        doc.close()
        return pages

    async def _optimize_image(self, image_path: str, temp_dir: str, page_num: int) -> str:
        """优化图片尺寸和质量"""
        def optimize_sync():
            try:
                img = Image.open(image_path)
                
                # 调整尺寸但保持宽高比
                img.thumbnail((settings.IMAGE_MAX_WIDTH, settings.IMAGE_MAX_HEIGHT), Image.Resampling.LANCZOS)
                
                # 保存优化后的图片
                optimized_path = os.path.join(temp_dir, f"optimized_page_{page_num}.png")
                img.save(optimized_path, "PNG", quality=settings.IMAGE_QUALITY, optimize=True)
                
                return optimized_path
            except Exception as e:
                logger.error(f"Image optimization error: {e}")
                return image_path  # 返回原图片路径
        
        loop = asyncio.get_event_loop()
        return await loop.run_in_executor(self.executor, optimize_sync)


    def _optimize_image_sync(self, image_path: str, temp_dir: str, page_num: int) -> str:
        """同步函数：实际执行图像优化"""
        try:
            img = Image.open(image_path)

            # 调整尺寸但保持宽高比
            img.thumbnail((settings.IMAGE_MAX_WIDTH, settings.IMAGE_MAX_HEIGHT), Image.Resampling.LANCZOS)

            # 保存优化后的图片
            optimized_path = os.path.join(temp_dir, f"optimized_page_{page_num}.png")
            img.save(optimized_path, "PNG", quality=settings.IMAGE_QUALITY, optimize=True)

            return optimized_path
        except Exception as e:
            logger.error(f"Image optimization error: {e}")
            return image_path  # 返回原图片路径